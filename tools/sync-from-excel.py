"""Read the English workbook and update the shared facts for all language views."""
import argparse,json,datetime,re,hashlib
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('workbook',nargs='?',type=Path,default=ROOT/'downloads/GoTracker_Catalog_EN.xlsx')
    parser.add_argument('--catalog',type=Path,default=ROOT/'data/catalog.json')
    args=parser.parse_args()
    data=json.loads(args.catalog.read_text(encoding='utf8'))
    wb=load_workbook(args.workbook,read_only=True,data_only=True)
    sheet=wb['Catalog'];rows=sheet.iter_rows(values_only=True)
    headers=next(rows)
    assert headers[:7]==('Dex No.','Image','Pokémon','In collection','ZIP folder','Released in GO','GO to HOME'),'Unexpected workbook columns'
    english=data['names']['en'];prefixes={'Alolan':'Alola','Galarian':'Galar','Hisuian':'Hisui','Paldean':'Paldea'}
    expected={e['id']:e for e in data['entries']};updates={}
    def boolean(value,nullable=False):
        text=str(value).strip()
        if nullable and text=='N/A':return None
        assert text in ('Yes','No'),f'Expected Yes/No, got {value!r}'
        return text=='Yes'
    for values in rows:
        if not values[0]:continue
        dex,_,name,available,folder,go,home=values[:7]
        dex=int(dex);region='Base';name=str(name).strip();form=None
        if dex==931:
            for color in ('green','blue','yellow','white'):
                suffix=f' ({color.title()} Plumage)'
                if name.endswith(suffix):name=name[:-len(suffix)];form=color;break
        for prefix,group in prefixes.items():
            if name.startswith(prefix+' '):region=group;name=name[len(prefix)+1:];break
        assert name==english[dex],f'Name and Dex do not match: {dex} {name}'
        key=f'{dex:04d}-{form if form and form!="green" else region}';assert key in expected and key not in updates,f'Unknown/duplicate row: {key}'
        present=boolean(available);go=boolean(go);home=boolean(home,True)
        folders=[] if folder is None or str(folder).strip() in ('','—','-') else sorted(set(re.split(r'\s*\n\s*',str(folder).strip())))
        assert present==bool(folders),f'Availability/folder mismatch: {key}'
        assert all(re.fullmatch(r'GoTracker/(SV|SS|BDSP|LZA|LGPE)',p) for p in folders),f'Invalid folder: {key}'
        assert go or home is None,f'GO/HOME mismatch: {key}'
        assert not present or go,f'File exists but GO is marked No: {key}'
        updates[key]={**expected[key],'available':present,'folders':folders,'go':go,'home':home}
    assert updates.keys()==expected.keys(),'Missing catalog entries in workbook'
    changed=sum(updates[k]!=expected[k] for k in expected)
    data['entries']=[updates[e['id']] for e in data['entries']]
    archive=ROOT/'downloads'/data['collectionFilename']
    checksum=hashlib.sha256(archive.read_bytes()).hexdigest() if archive.exists() else data['collectionSha256']
    if changed or checksum!=data['collectionSha256']:data['updated']=datetime.date.today().isoformat()
    data['collectionSha256']=checksum
    args.catalog.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    print(f'{len(updates)} rows synchronized; {changed} changed. All languages use these shared facts.')
if __name__=='__main__':main()
